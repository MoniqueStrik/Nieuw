"""
Build Coda tab in Format_2026_begrotingswijziging_raad.xlsx
Spec: PROMPT_Coda_tab_Format2026_1.md
"""
import copy
import openpyxl
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.utils import get_column_letter

TARGET_PATH = '/root/.claude/uploads/f57bf805-679c-56f2-b85d-87735f25bc3e/a9123695-Format_2026_begrotingswijziging_raad_inlees_Coda.xlsx'
CODA_SRC_PATH = '/root/.claude/uploads/f57bf805-679c-56f2-b85d-87735f25bc3e/319558b9-Inleesformat_Coda.xlsx'
OUTPUT_PATH = '/home/user/Nieuw/Format_2026_begrotingswijziging_raad.xlsx'

# ─── helpers ──────────────────────────────────────────────────────────────────

def copy_cell_style(src, dst):
    dst.font = copy.copy(src.font)
    dst.fill = copy.copy(src.fill)
    dst.border = copy.copy(src.border)
    dst.alignment = copy.copy(src.alignment)
    dst.number_format = src.number_format

def copy_cell(src, dst):
    dst.value = src.value
    copy_cell_style(src, dst)

def clear_cell(cell):
    cell.value = None
    cell.font = Font()
    cell.fill = PatternFill()
    cell.border = Border()
    cell.alignment = Alignment()
    cell.number_format = 'General'

def IG(col, row):
    """Reference to 'Invullen groen'!{col}{row}"""
    return f"'Invullen groen'!{col}{row}"

IG_NAME = "'Invullen groen'"

# ─── omschrijving formula building ────────────────────────────────────────────

def tekstbron(n):
    """§3 tekstbron: pick G or K from 'Invullen groen' row n."""
    return (
        f'IF(LEFT(TEXT({IG_NAME}!E{n},"0"),1)="6",'
        f'{IG_NAME}!K{n},'
        f'IF({IG_NAME}!J{n}>0,{IG_NAME}!G{n},{IG_NAME}!K{n}))'
    )

def clean_formula(inner):
    """Apply cleaning steps to an Excel text formula (already inside LOWER)."""
    # years
    result = inner
    for year in ('2024', '2025', '2026', '2027', '2028', '2029'):
        result = f'SUBSTITUTE({result},"{year}","")'
    # taakveld
    result = f'SUBSTITUTE({result},"taakveld ","")'
    # abbreviations
    abbrevs = [
        ('installaties', 'inst'),
        ('ruwbouw', 'ruwb'),
        ('afbouw', 'afb'),
        ('bijstelling', 'bijst'),
        ('investeringen', 'invest'),
    ]
    for long, short in abbrevs:
        result = f'SUBSTITUTE({result},"{long}","{short}")'
    return f'TRIM({result})'

def raw_text_formula_61(n):
    """Helper col J for section 6.1: only show when E{n} (grootboek) filled."""
    tb = tekstbron(n)
    cleaned = clean_formula(f'LOWER({tb})')
    return (
        f'=IF({IG_NAME}!E{n}>0,'
        f'IFERROR($K$1&" "&{cleaned},""),'
        f'"")'
    )

def raw_text_formula_62(n):
    """Helper col J for section 6.2: only show when J{n} filled and J{n}<>2999999."""
    tb = tekstbron(n)
    cleaned = clean_formula(f'LOWER({tb})')
    return (
        f'=IF(AND({IG_NAME}!J{n}>0,{IG_NAME}!J{n}<>2999999),'
        f'IFERROR($K$1&" "&{cleaned},""),'
        f'"")'
    )

def E_omschrijving_formula(coda_row):
    """Word-boundary truncation at 36 chars, referencing J{coda_row}."""
    r = f'J{coda_row}'
    last_space = (
        f'FIND(CHAR(1),SUBSTITUTE(LEFT({r},36)," ",CHAR(1),'
        f'LEN(LEFT({r},36))-LEN(SUBSTITUTE(LEFT({r},36)," ",""))))-1'
    )
    return (
        f'=IFERROR('
        f'IF(LEN({r})<=36,{r},'
        f'IFERROR(LEFT({r},{last_space}),LEFT({r},36))),'
        f'"")'
    )

def E_sluitrekening_formula():
    return '=$K$1&" sluitrekening balans"'

# ─── amount formula ───────────────────────────────────────────────────────────

def amount_formula(ig_col, n):
    ref = f"{IG_NAME}!{ig_col}{n}"
    return f'=IFERROR(IF({ref}=0,"",INT({ref})),"")'

# ─── main ─────────────────────────────────────────────────────────────────────

def main():
    print("Laden werkboeken...")
    wb = openpyxl.load_workbook(TARGET_PATH, data_only=False)
    wb_src = openpyxl.load_workbook(CODA_SRC_PATH, data_only=False)

    ws_src = wb_src['Format 2026 begr.wijzigingen']
    ws = wb['Coda']

    # ── Stap 1: Koptekst rijen 1-4 overnemen ──────────────────────────────────
    print("Stap 1: koptekst kopiëren...")

    # First clear all existing content in Coda from row 1
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            clear_cell(cell)

    # Copy rows 1-4 cell by cell from source
    for r in range(1, 5):
        for c in range(1, ws_src.max_column + 1):
            src_cell = ws_src.cell(row=r, column=c)
            dst_cell = ws.cell(row=r, column=c)
            copy_cell(src_cell, dst_cell)

    # Copy column widths
    for col_letter, dim in ws_src.column_dimensions.items():
        if dim.width:
            ws.column_dimensions[col_letter].width = dim.width

    # Override C1: formula for begrotingswijzigingsnummer
    ws['C1'].value = "='Invullen groen'!L4"

    # ── Helper K1: prefix ──────────────────────────────────────────────────────
    # prefix: "W12/2026" → "W12-26"
    ws['K1'].value = (
        '=IFERROR('
        'LEFT(\'Invullen groen\'!L4,FIND("/",\'Invullen groen\'!L4)-1)'
        '&"-"&'
        'RIGHT(\'Invullen groen\'!L4,2),'
        '"")'
    )
    ws.column_dimensions['K'].hidden = True
    ws.column_dimensions['J'].hidden = True

    # ── Stap 2-4: Data rijen ───────────────────────────────────────────────────
    print("Stap 2-4: data rijen schrijven...")

    # Row mapping
    sections_61 = [(17, 74), (78, 109)]
    sections_62 = [(121, 167), (171, 216), (220, 272)]

    coda_row = 5

    # ── Sectie 6.1: exploitatie ────────────────────────────────────────────────
    for ig_start, ig_end in sections_61:
        for n in range(ig_start, ig_end + 1):
            # A: grootboek
            ws.cell(row=coda_row, column=1).value = (
                f'=IFERROR(IF({IG_NAME}!E{n}>0,{IG_NAME}!E{n},""),"")'
            )
            # B: econ.categorie
            ws.cell(row=coda_row, column=2).value = (
                f'=IFERROR(IF({IG_NAME}!H{n}>0,{IG_NAME}!H{n},""),"")'
            )
            # C: verbijzondering
            ws.cell(row=coda_row, column=3).value = (
                f'=IFERROR(IF({IG_NAME}!J{n}>0,{IG_NAME}!J{n},""),"")'
            )
            # D: leeg
            ws.cell(row=coda_row, column=4).value = None
            # J: helper raw text (alleen tonen als rij gevuld)
            ws.cell(row=coda_row, column=10).value = raw_text_formula_61(n)
            # E: omschrijving (word-boundary truncation)
            ws.cell(row=coda_row, column=5).value = E_omschrijving_formula(coda_row)
            # F-I: bedragen
            for offset, ig_col in enumerate(['L', 'M', 'N', 'O'], start=6):
                c = ws.cell(row=coda_row, column=offset)
                c.value = amount_formula(ig_col, n)
                c.number_format = '0'
            coda_row += 1

    balans_start_coda = coda_row  # eerste balansrij in Coda

    # ── Sectie 6.2: balans ─────────────────────────────────────────────────────
    for ig_start, ig_end in sections_62:
        for n in range(ig_start, ig_end + 1):
            # A: "+" als J gevuld en J <> 2999999 (gecentreerd)
            a_cell = ws.cell(row=coda_row, column=1)
            a_cell.value = (
                f'=IF(AND({IG_NAME}!J{n}>0,{IG_NAME}!J{n}<>2999999),"+","")'
            )
            a_cell.alignment = Alignment(horizontal='center')
            # B: econ.categorie (zal leeg zijn voor balansrijen, spec §2)
            ws.cell(row=coda_row, column=2).value = (
                f'=IFERROR(IF({IG_NAME}!H{n}>0,{IG_NAME}!H{n},""),"")'
            )
            # C: verbijzondering
            ws.cell(row=coda_row, column=3).value = (
                f'=IFERROR(IF({IG_NAME}!J{n}>0,{IG_NAME}!J{n},""),"")'
            )
            # D: leeg
            ws.cell(row=coda_row, column=4).value = None
            # J: helper raw text (alleen tonen als rij gevuld)
            ws.cell(row=coda_row, column=10).value = raw_text_formula_62(n)
            # E: omschrijving
            ws.cell(row=coda_row, column=5).value = E_omschrijving_formula(coda_row)
            # F-I: bedragen
            for offset, ig_col in enumerate(['L', 'M', 'N', 'O'], start=6):
                c = ws.cell(row=coda_row, column=offset)
                c.value = amount_formula(ig_col, n)
                c.number_format = '0'
            coda_row += 1

    balans_end_coda = coda_row - 1  # laatste balansrij in Coda
    sluitrekening_row = coda_row

    # ── Sectie 6.3: sluitrekening balans ──────────────────────────────────────
    print(f"Stap 4.4: sluitrekening op rij {sluitrekening_row}...")
    ws.cell(row=sluitrekening_row, column=1).value = 2999999
    ws.cell(row=sluitrekening_row, column=2).value = 86100
    ws.cell(row=sluitrekening_row, column=3).value = None
    ws.cell(row=sluitrekening_row, column=4).value = None
    ws.cell(row=sluitrekening_row, column=5).value = E_sluitrekening_formula()

    for offset, col_letter in enumerate(['F', 'G', 'H', 'I'], start=6):
        c = ws.cell(row=sluitrekening_row, column=offset)
        c.value = (
            f'=IFERROR(-SUM({col_letter}{balans_start_coda}:{col_letter}{balans_end_coda}),0)'
        )
        c.number_format = '0'

    # ── Stap 7: Verwijder gemeente Best-logo (geen drawing gevonden) ───────────
    if ws._drawing is not None:
        print("Waarschuwing: drawing aanwezig in Coda-tabblad, handmatig verwijderen vereist.")
    else:
        print("Geen logo/drawing gevonden in Coda-tabblad.")

    # ── Opslaan ────────────────────────────────────────────────────────────────
    print(f"Opslaan naar {OUTPUT_PATH}...")
    wb.save(OUTPUT_PATH)
    print(f"Klaar. Coda bevat {sluitrekening_row} rijen ({sluitrekening_row-4} datarijen).")
    print(f"  Sectie 6.1: Coda rijen 5-{balans_start_coda-1}")
    print(f"  Sectie 6.2: Coda rijen {balans_start_coda}-{balans_end_coda}")
    print(f"  Sectie 6.3: Coda rij {sluitrekening_row}")

    return {
        'balans_start': balans_start_coda,
        'balans_end': balans_end_coda,
        'sluitrekening': sluitrekening_row,
    }


if __name__ == '__main__':
    result = main()
