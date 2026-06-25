"""
Fix ontbrekende VLOOKUP-formules in 'Invullen groen'.

Voegt alleen toe aan cellen die momenteel leeg zijn (None).
Overschrijft niets. Wijzigt geen andere tabbladen.

Ontbrekende formules per sectie:
  Sectie A (rijen 17-74):
    - Rij 17: A, B, C, D, G, H
    - Rij 18: A, H
    - Rijen 19-74: H
  Sectie B (rijen 78-109):
    - Alle rijen: B, C, G, H
"""
import openpyxl

FILE = '/home/user/Nieuw/Format_2026_begrotingswijziging_raad.xlsx'

# Formule-patronen (matchen bestaande formules in het bestand)
def f_A(n): return f'=IFERROR(VLOOKUP(E{n},Tabellen!A:F,4),0)'   # omschrijving programma
def f_B(n): return f'=IFERROR(VLOOKUP(E{n},Tabellen!A:F,3),0)'   # nr hoofdtaakveld (htv)
def f_C(n): return f'=IFERROR(VLOOKUP(E{n},Tabellen!A:F,5),0)'   # nr taakveld
def f_D(n): return f'=IFERROR(VLOOKUP(E{n},Tabellen!A:F,6),0)'   # omschrijving taakveld
def f_G(n): return f'=IFERROR(VLOOKUP(E{n},Tabellen!A:F,2),0)'   # omschrijving grootboek
def f_H(n): return f'=IFERROR(VLOOKUP(E{n},Tabellen!A:H,8),0)'   # nr kostencategorie (Cat)

def fill_if_empty(ws, row, col, formula):
    """Vul cel alleen als die momenteel leeg is."""
    cell = ws.cell(row=row, column=col)
    if cell.value is None:
        cell.value = formula
        return True
    return False

def main():
    wb = openpyxl.load_workbook(FILE, data_only=False)
    ws = wb['Invullen groen']

    added = 0

    # ── Sectie A: rijen 17-74 ─────────────────────────────────────────────────
    for n in range(17, 75):
        if fill_if_empty(ws, n, 1, f_A(n)): added += 1   # col A
        if fill_if_empty(ws, n, 2, f_B(n)): added += 1   # col B
        if fill_if_empty(ws, n, 3, f_C(n)): added += 1   # col C
        if fill_if_empty(ws, n, 4, f_D(n)): added += 1   # col D
        if fill_if_empty(ws, n, 7, f_G(n)): added += 1   # col G
        if fill_if_empty(ws, n, 8, f_H(n)): added += 1   # col H

    # ── Sectie B: rijen 78-109 ────────────────────────────────────────────────
    for n in range(78, 110):
        if fill_if_empty(ws, n, 1, f_A(n)): added += 1   # col A
        if fill_if_empty(ws, n, 2, f_B(n)): added += 1   # col B
        if fill_if_empty(ws, n, 3, f_C(n)): added += 1   # col C
        if fill_if_empty(ws, n, 4, f_D(n)): added += 1   # col D
        if fill_if_empty(ws, n, 7, f_G(n)): added += 1   # col G
        if fill_if_empty(ws, n, 8, f_H(n)): added += 1   # col H

    print(f'{added} formules toegevoegd aan Invullen groen.')

    # Verificatie: controleer een paar rijen
    print('\nVerificatie:')
    for r in [17, 18, 19, 78]:
        row_vals = {
            'A': ws.cell(r, 1).value,
            'B': ws.cell(r, 2).value,
            'C': ws.cell(r, 3).value,
            'D': ws.cell(r, 4).value,
            'G': ws.cell(r, 7).value,
            'H': ws.cell(r, 8).value,
        }
        missing = [k for k, v in row_vals.items() if v is None]
        status = '✓' if not missing else f'✗ mist: {missing}'
        print(f'  Rij {r}: {status}')

    wb.save(FILE)
    print('\nOpgeslagen.')

if __name__ == '__main__':
    main()
